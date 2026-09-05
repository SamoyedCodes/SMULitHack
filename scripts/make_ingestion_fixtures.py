"""Create public synthetic ingestion fixtures only; no extraction answer key implied."""
import argparse
import io
from pathlib import Path

import pymupdf
from PIL import Image, ImageEnhance, ImageFilter
from docx import Document

TEXT = '1. Synthetic agreement\nMeridian Test Pte Ltd and Example Supplier Pte Ltd.\nThis synthetic document is for local ingestion testing only.\n2. Notice\nWritten notice must be received sixty calendar days before expiry.'


def native_pdf(text=TEXT, pages=1):
    pdf = pymupdf.open()
    for i in range(pages):
        page = pdf.new_page(width=612, height=792)
        page.insert_textbox(pymupdf.Rect(45, 60, 565, 680), f'{text}\nPage {i+1}', fontsize=16)
    data = pdf.tobytes()
    pdf.close()
    return data


def fixtures():
    native = native_pdf()
    with pymupdf.open(stream=native, filetype='pdf') as pdf:
        png = pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).tobytes('png')
    def scanned(image):
        pdf = pymupdf.open(); page = pdf.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=image)
        data = pdf.tobytes(); pdf.close(); return data
    image = Image.open(io.BytesIO(png)).convert('RGB')
    degraded = ImageEnhance.Contrast(image.filter(ImageFilter.GaussianBlur(1.1))).enhance(.45)
    buffer = io.BytesIO(); degraded.save(buffer, 'PNG')
    jpeg = io.BytesIO(); image.save(jpeg, 'JPEG')
    mixed = pymupdf.open(stream=native, filetype='pdf')
    page = mixed.new_page(width=612, height=792); page.insert_image(page.rect, stream=png)
    mixed.new_page(width=612, height=792)  # Deliberately blank: visibly unresolved reading.
    docx = Document(); docx.add_heading('Synthetic signed agreement fixture', 0); docx.add_paragraph(TEXT)
    word = io.BytesIO(); docx.save(word)
    result = {'native.pdf': native, 'clean-scan.pdf': scanned(png), 'degraded-scan.pdf': scanned(buffer.getvalue()),
              'mixed-pages.pdf': mixed.tobytes(), 'scan.png': png, 'scan.jpeg': jpeg.getvalue(), 'agreement.docx': word.getvalue()}
    mixed.close()
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--output', type=Path, required=True)
    output = parser.parse_args().output; output.mkdir(parents=True, exist_ok=True)
    for name, data in fixtures().items():
        (output / name).write_bytes(data)
    print(f'Created 7 synthetic ingestion files in {output}')
